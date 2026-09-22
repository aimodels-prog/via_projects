"""Export demo values into the existing six-column PDF template, without new fields."""
from pathlib import Path
import csv
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import shutil
import subprocess

root=Path(__file__).resolve().parents[1]
demo=load_workbook(root/'deliverables/VIA_Mughsayl_Demo_Current_Template.xlsx',data_only=True)
values={r[0]:r[2] for name in ['Project','Contract_Dates','Progress_Finance'] for r in demo[name].iter_rows(min_row=4,values_only=True)}
mapping={
 'Project number':'project_reference','Project type':'project_type','Report number':'report_number',
 'Project name':'project_name','Region':'region','Approved contract total':'dashboard_contract_value',
 'Currency':'currency','Award date':'award_date','Start date':'start_date','Completion date':'completion_date',
 'Mobilization days':'mobilization_days','Construction days':'construction_days',
 'Construction days with VOs':'construction_days_with_vos','Completion date with VOs':'completion_date_with_vos',
 'Brief description':'brief','Data as of':'data_as_of','Revision':'revision','Elapsed days':'elapsed_days',
 'Remaining days':'remaining_days','Expected completion date':'expected_completion',
 'Cumulative planned progress %':'physical_planned','Cumulative actual progress %':'physical_actual',
 'This month planned %':'monthly_planned','This month actual %':'monthly_actual',
 'Planned financial progress %':'financial_planned','Actual financial progress %':'financial_actual',
 'Actually paid amount':'paid_amount',
}
direct={label:values.get(key) for label,key in mapping.items()}
direct['Footer code']='DUMMY DATA — NOT FOR REPORTING'
for row,label in [(4,'Client'),(5,'Consultant'),(6,'Contractor')]:
    ws=demo['Contacts']
    direct[label]=ws.cell(row,2).value
    if label=='Client':direct['Client department']=ws.cell(row,3).value
    else:
        for suffix,col in [('representative',4),('role',5),('phone',6)]:direct[f'{label} {suffix}']=ws.cell(row,col).value
direct['Client engineer']=demo['Contacts']['D7'].value
direct['Client engineer phone']=demo['Contacts']['F7'].value
for row,label in [(4,'machinery'),(5,'manpower')]:
    direct[f'Planned {label}']=demo['Resources'].cell(row,2).value
    direct[f'Actual {label}']=demo['Resources'].cell(row,3).value

with (root/'public/templates/pdf-report-template.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
activities=[r for r in demo['Activities'].iter_rows(min_row=4,values_only=True) if r[1]]
assets={r[0]:r[4] for r in demo['Assets'].iter_rows(min_row=4,values_only=True)}
def cell(v):
    if v is None:return ''
    if isinstance(v,datetime):return v.strftime('%Y-%m-%d')
    return str(v)
i=0
schedule=iter([r for r in demo['Monthly_Schedule'].iter_rows(min_row=4,values_only=True) if r[0]])
scopes=iter([r for r in demo['Scope_Layers'].iter_rows(min_row=4,values_only=True) if r[0]=='Scope'])
layers=iter([r for r in demo['Scope_Layers'].iter_rows(min_row=4,values_only=True) if r[0]=='Layer'])
trades=iter([r for r in demo['Status_Updates'].iter_rows(min_row=4,values_only=True) if r[1]])
for row in rows[1:]:
    section,label=row[:2]
    if section in ['project','monthly']:row[2]=cell(direct.get(label))
    elif section=='activity' and i<len(activities):
        a=activities[i];i+=1;row[1]=a[1];row[3]=cell(a[2]);row[4]=cell(a[3])
    elif section=='photo':row[2]=assets[label] or ''
    elif section=='layout':row[2]='raysut-reference'
    elif section=='schedule':
        s=next(schedule,None)
        if s:row[1]=s[0].strftime('%b-%y');row[3]=cell(s[1]);row[4]=cell(s[2])
    elif section=='scope':
        s=next(scopes,None)
        if s:row[1]=s[2];row[2]=cell(s[3]);row[3]=cell(s[4])
    elif section=='layer':
        s=next(layers,None)
        if s:row[1]=s[1];row[2]=s[2];row[3]=cell(s[5])
    elif section=='trade':
        s=next(trades,None)
        if s:row[1]=s[1];row[2]=s[2].lower().replace(' ','')
    row[5]='DUMMY DEMO DATA — NOT FOR REPORTING. '+row[5]
    if section=='contract':row[5]+=' Demo does not specify this contract breakdown; left blank.'
name='pdf-report-template-Mughsayl-Demo'
out=root/'deliverables'/f'{name}.csv'
with out.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
wb=Workbook();wb.remove(wb.active)
tabs=[('01 Project setup',['project']),('02 Monthly update',['monthly']),('03 Contract values',['contract']),('04 Activities',['activity']),('05 S-curve schedule',['schedule']),('06 Scope quantities',['scope']),('07 Pavement layers',['layer']),('08 Trade statuses',['trade']),('09 Photo captions',['photo']),('10 Project layout',['layout','layout_section'])]
for title,sections in tabs:
    ws=wb.create_sheet(title);ws.append(rows[0])
    for row in rows[1:]:
        if row[0] in sections:ws.append(row)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    ws.sheet_properties.tabColor='005A9C'
    ws.column_dimensions['A'].hidden=True
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='15354A');c.font=Font(color='FFFFFF',bold=True)
    for col,width in [('A',14),('B',38),('C',65),('D',14),('E',14),('F',95)]:ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row[0].row].height=44
wb.save(root/'deliverables'/f'{name}.xlsx')
for suffix in ['csv','xlsx']:shutil.copy2(root/'deliverables'/f'{name}.{suffix}',root/'public/demos'/f'{name}.{suffix}')
assert all(len(r)==6 for r in rows)
assert rows[0]==['section','field','value','planned','actual','instructions']
print('Created CSV and sectioned Excel using the PDF template. No project data changed.')
subprocess.run(['npx.cmd','tsx','scripts/sync-pdf-template.ts'],cwd=root,check=True)
