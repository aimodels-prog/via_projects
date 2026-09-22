"""Generate versioned input workbooks, not automatic source-verification results.
Requires: openpyxl. Run from the project root. Generated files live in deliverables/.
"""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

OUT = Path(__file__).resolve().parents[1] / 'deliverables'
BLUE, NAVY, GREEN = 'EAF3FF', '15354A', 'E8F3E9'
SOURCE = 'Sample PDF: page 1 (printed page 4)'
STATUS = '"Not reported,Not applicable,Entered,Needs review,Verified"'

# key, display label, unit/type, requirement, instructions, sample value
PROJECT = [
 ('project_name','Full project name','text','Required','Copy the approved project title.','CONSTRUCTION WORKS OF MUGHSAYL ROAD AND BRIDGE'),
 ('project_reference','Project / contract reference','text','If reported','Do not invent an identifier.',None),
 ('project_slug','Existing project URL slug','text','Required','For later months use the SAME project slug. Set once in the app.',None),
 ('project_type','Project type','text','Required','Internal classification; confirm with project team.',None),
 ('region','Region / governorate','text','Required','As printed.','Governorate of Dhofar'),
 ('report_month','Reporting month','month','Required','Enter first day of the report month; displayed as month and year.',datetime(2026,7,1)),
 ('report_number','Report number','text','If reported','Printed page number is NOT the report number.',None),
 ('data_as_of','Exact data-as-of date','date','Required','Confirm exact date; July 2026 alone does not establish 31 July.',None),
 ('revision','Report revision','text','If reported','Do not assume Rev-01.',None),
 ('source_filename','Original report filename','text','Required','Source PDF retained for review.','Extracted pages from JULY- Mughsayl Bridge and Road Works.pdf'),
 ('brief','Project brief','text','Required','Keep full scope and significant qualifications.','Mughsayl Bridge & Road Works; Mughsayl Bridge – 540m; Marnif cave access road 2 km approx.'),
 ('prepared_by','Prepared by','text','Internal','Name of person entering this workbook.',None),
 ('checked_by','Checked by','text','Internal','Reviewer name. A workbook name is not an authenticated app approval.',None),
 ('checked_date','Checked date','date','Internal','Actual date reviewed.',None),
]
CONTRACT = [
 ('currency','Currency','text','Required','Applies to monetary figures.','R.O.'),
 ('original_ex_contingency','Original contract excluding contingency','amount','Required','Preserve three decimal places.',9110540.534),
 ('contingency','Contingency','amount','If reported','Sample prints Nil; numeric zero is its interpretation, subject to review.',0),
 ('original_inc_contingency','Original contract including contingency','amount','Required','Do not sum variation totals together.',9110540.534),
 *[(f'contract_with_vo_{i}',f'Contract with VO.{i} and contingency','amount','If applicable','Sample prints NA: keep value blank and select Not applicable.',None) for i in range(1,5)],
 ('dashboard_contract_value','Approved contract value used on dashboard','amount','Required','Reviewer selects the applicable current contract basis; do not guess.',None),
 ('contract_basis','Basis of dashboard contract value','text','Required','For example: original including contingency, or approved VO.2 total.',None),
 ('award_date','Awarding date','date','Required','Enter a real Excel date.',datetime(2024,5,8)),
 ('mobilization_days','Mobilization period','days','If reported','Not automatically added to construction duration.',60),
 ('construction_days','Original construction period','days','Required','Sample: 669 days.',669),
 ('start_date','Start of construction','date','Required','As printed; reconcile with duration separately.',datetime(2024,11,2)),
 ('completion_date','Contract completion date','date','Required','Original contractual completion date.',datetime(2026,9,2)),
 ('construction_days_with_vos','Construction period with VOs','days','If applicable','Sample prints a dash; do not interpret this as zero.',None),
 ('completion_date_with_vos','Completion date with VOs','date','If applicable','Leave blank if unreported.',None),
 ('elapsed_days','Elapsed time','days','Required','Reported duration; check against remaining days.',636),
 ('remaining_days','Remaining period','days','Required','Do not calculate from today; this is a monthly snapshot.',33),
 ('expected_completion','Expected completion date','date','Required','Separate from original completion date.',datetime(2026,9,2)),
]
PROGRESS = [
 ('physical_planned','Bridge & road cumulative planned progress','percent','Required','Enter 95 for 95%, NOT 0.95.',95),
 ('physical_actual','Bridge & road cumulative actual progress','percent','Required','Do not substitute financial progress.',98.45),
 ('physical_variance_reported','Reported physical ahead / delay','signed percent','If reported','Keep printed sign; calculated actual-minus-planned appears in Checks.',3.45),
 ('arches_planned','Arches cumulative planned progress','percent','If reported','Arches column is blank in sample. Activity-row figures belong on Activities.',None),
 ('arches_actual','Arches cumulative actual progress','percent','If reported','Do not copy overall bridge progress into this column.',None),
 ('monthly_planned','This-month planned physical progress','percent','If reported','Leave blank if tiny chart table is unreadable.',None),
 ('monthly_actual','This-month actual physical progress','percent','If reported','Never derive by estimating line heights.',None),
 ('financial_planned','Financial cumulative planned progress','percent','Required','Independent of physical planned progress.',85.91),
 ('financial_actual','Financial cumulative actual progress','percent','Required','Confirm what this measure represents; do not assume paid / contract.',81.52),
 ('financial_difference_reported','Printed financial difference','signed percent','If reported','Sample prints +4.39 without a minus; actual-minus-planned is -4.39.',4.39),
 ('financial_difference_basis','Meaning of reported financial difference','text','Needs clarification','Confirm whether the source uses planned minus actual.',None),
 ('anticipated_payment','Total payment anticipated','amount','If reported','This is NOT total paid or certified amount.',400000),
 ('paid_amount','Total actually paid','amount','If reported','Not explicitly provided in sample. Keep blank.',None),
 ('certified_amount','Total certified amount','amount','If reported','Distinct from anticipated and actually paid amounts.',None),
 ('financial_basis','Definition / denominator of financial percentage','text','Needs clarification','Use to determine whether a paid-to-contract comparison is meaningful.',None),
 ('chart_mode','S-curve representation','text','Required','Choose Verified schedule or Source image; sample needs original chart/table.',None),
]

def sheet(wb, name, title, note, headers, rows):
    ws=wb.create_sheet(name)
    ws.append([title]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws.append([note]);ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
    ws.append(headers)
    for row in rows: ws.append(row)
    ws.freeze_panes='C4';ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=30;ws.row_dimensions[2].height=48;ws.row_dimensions[3].height=32
    for cell in ws[1]+ws[3]: cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(color='FFFFFF',bold=True,size=12 if cell.row==1 else 10)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    for row in ws.iter_rows(min_row=4):
        ws.row_dimensions[row[0].row].height=46
        for c in row: c.font=Font(name='Calibri',size=11,color='16324F');c.fill=PatternFill('solid',fgColor=BLUE)
    for i in range(1,len(headers)+1): ws.column_dimensions[ws.cell(3,i).column_letter].width=24
    if ws.max_row>3:
        table=Table(displayName='T_'+name.replace(' ','_'),ref=f'A3:{ws.cell(ws.max_row,len(headers)).coordinate}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:3'
    return ws

def dropdown(ws, cells, formula):
    dv=DataValidation(type='list',formula1=formula,allow_blank=True);dv.errorTitle='Choose a listed value';dv.error='Use the dropdown choices.';dv.showErrorMessage=True
    ws.add_data_validation(dv);dv.add(cells)

def numeric(ws,cells,low=0,high=1000000000000):
    dv=DataValidation(type='decimal',operator='between',formula1=str(low),formula2=str(high),allow_blank=True)
    dv.showErrorMessage=True;dv.error=f'Enter a number between {low} and {high}; leave missing values blank.'
    ws.add_data_validation(dv);dv.add(cells)

def fields(wb,name,items,example):
    rows=[]
    for key,label,unit,required,note,value in items:
        status='Needs review' if example and value is not None else 'Not reported'
        if example and key.startswith('contract_with_vo_'):status='Not applicable'
        rows.append([key,label,value if example else None,unit,required,status,SOURCE if example else '',note])
    ws=sheet(wb,name,name.replace('_',' '),'Blue cells: enter data. Keep field keys, headers and sheet names unchanged. Missing numeric data = blank + a status, never invented zero.',
             ['Field key','Field label','Value','Unit / type','Requirement','Entry status','Source reference','Instructions / clarification'],rows)
    ws.column_dimensions['A'].width=33;ws.column_dimensions['B'].width=43;ws.column_dimensions['C'].width=40;ws.column_dimensions['H'].width=64;ws.column_dimensions['G'].width=35
    refs={}
    for r,item in enumerate(items,4):
        key,label,unit,*_=item;refs[key]=f"'{name}'!C{r}"
        ws.cell(r,1).fill=PatternFill('solid',fgColor='F1F3F5')
        if unit in ('amount','days','percent','signed percent'):
            numeric(ws,f'C{r}',-100 if unit=='signed percent' else 0,100 if 'percent' in unit else 1000000000000)
            ws.cell(r,3).number_format='0.00"%"' if 'percent' in unit else '#,##0.000' if unit=='amount' else '0'
        elif unit in ('date','month'):ws.cell(r,3).number_format='mmm yyyy' if unit=='month' else 'yyyy-mm-dd'
        else:ws.cell(r,3).number_format='@'
    dropdown(ws,f'F4:F{ws.max_row}',STATUS)
    return refs

def build(example=False):
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    guidance=[
      ['Template version','VIA-REPORT-EXCEL-1.0; one workbook = one project + one reporting month.'],
      ['Workbook purpose','Data-entry specification based on Mughsayl summary PDF, NOT an approved report.'],
      ['Current app compatibility','Direct .xlsx import still needs implementation. Do not upload this workbook to the existing CSV-only importer. This workbook defines the proposed Excel import contract.'],
      ['Start here','Complete Project, Contract_Dates and Progress_Finance. Then complete relevant repeating tables.'],
      ['Numbers','Enter 95 for 95 percent; percent cells display 95.00%. Enter amounts as numbers, no currency text. Use zero only when explicitly reported.'],
      ['Missing vs not applicable','Leave missing values blank and select Not reported. Use Not applicable only when the source/team confirms it. Never put NA into a numeric cell.'],
      ['Dates / identifiers','Use real Excel dates. Report month uses its first day only as a month identifier. Store references and telephone numbers as text to preserve leading zeroes.'],
      ['Manual assets','Upload logo, project layout and four photographs in the app; record captions and asset IDs on Assets. Additional logos, status drawing and source S-curve image have separate slots. No images need to be embedded in Excel.'],
      ['Schedule','Copy readable monthly numbers from original Excel/source table. Leave future actuals blank. If unreadable, choose Source image and upload the original S-curve. Never estimate plotted values.'],
      ['Financial semantics','Financial progress, certified amount, paid amount and anticipated payment are distinct. Never substitute one for another.'],
      ['Scope and layers','Optional if not reported / not applicable. Never invent five quantities or pavement layers to fill a dashboard.'],
      ['Editing tables','Keep sheet names, field keys and table headers. Fill existing rows; add rows inside the Excel table if needed. Ignore unused empty rows on import.'],
      ['Checks','Formula cells are green. Excel recalculates them on opening; they are review aids only. The application must independently validate values on import.'],
      ['Review','Needs review is NOT verified. Check against the source, mark verified fields, then preview and approve in the app. Workbook approval names do not replace secure app approval.'],
      ['Example limitations','Example figures are manually transcribed from the supplied image-based PDF and require checking. Tiny S-curve values, unclear contacts and absent data are intentionally left blank.'],
      ['Sample issue to resolve','Financial difference is printed 4.39%; actual minus planned is -4.39 percentage points. Preserve both and confirm the source convention.'],
      ['Source coverage','Includes contract/VOs, dates, bridge and arches progress, financial figures, activity table, monthly schedule, resources, scope, narrative updates, contacts, photographs, logos and the separate activity-status drawing.'],
    ]
    guide=sheet(wb,'Read_Me','VIA monthly project reporting | '+('UNVERIFIED MUGHSAYL EXAMPLE' if example else 'BLANK INPUT TEMPLATE'),'Do not publish without source review. This workbook contains no client passwords.',['Topic','Instructions'],guidance)
    guide.column_dimensions['A'].width=28;guide.column_dimensions['B'].width=125
    refs={};refs.update(fields(wb,'Project',PROJECT,example));refs.update(fields(wb,'Contract_Dates',CONTRACT,example));refs.update(fields(wb,'Progress_Finance',PROGRESS,example))
    dropdown(wb['Progress_Finance'],refs['chart_mode'].split('!')[1],'"Verified schedule,Source image"')
    acts=[('4 Cell Box Culvert',100,100,0),('Animal Crossing',100,100,0),('Earthwork and Subgrade (incl. parking area)',100,85,-15),('Bridge Substructure',100,100,0),('Bridge Superstructure',100,100,0),('Arches Works',50,50,0),('Slope Protection Works',100,65,-35),('Lightpole Fixtures',100,100,0),('Finishing Works (Road Markings + Sign Boards)',85,90,5)]
    rows=[]
    for i in range(20):
        a=acts[i] if example and i<len(acts) else (None,None,None,None);r=i+4
        rows.append([f'activity_{i+1:02}',*a,f'=IF(COUNT(C{r}:D{r})=2,D{r}-C{r},"")','Needs review' if a[0] else 'Not reported',SOURCE if a[0] else '',None])
    ws=sheet(wb,'Activities','Physical progress of major activities','Percentages are percentage points: 100 = 100%. Above 100 is retained for review, not silently capped.',['Activity ID','Activity name','Planned %','Actual %','Printed difference %','Calculated actual-plan','Entry status','Source reference','Notes'],rows)
    ws.column_dimensions['B'].width=48;dropdown(ws,'G4:G23',STATUS);numeric(ws,'C4:D23',0,1000);numeric(ws,'E4:E23',-1000,1000)
    rows=[]
    for r in range(4,64):rows.append([None,None,None,None,None,f'=IF(COUNT(B{r}:C{r})=2,C{r}-B{r},"")',None,None,None])
    ws=sheet(wb,'Monthly_Schedule','S-curve | monthly source values','Leave month cells blank until original values are available. Chronological order, one row per month; no invented zero start point.',['Month','Monthly planned %','Monthly actual %','Cumulative planned %','Cumulative actual %','Monthly actual-plan','Entry status','Source reference','Notes'],rows)
    dropdown(ws,'G4:G63',STATUS);numeric(ws,'B4:E63',0,100)
    for r in range(4,64):ws.cell(r,1).number_format='mmm yyyy'
    rows=[]
    for i,(name,p,a) in enumerate([('Machinery',40,35),('Manpower',70,65)]):
        r=i+4;rows.append([name,p if example else None,a if example else None,f'=IF(COUNT(B{r}:C{r})=2,C{r}-B{r},"")',-5 if example else None,'units' if i==0 else 'persons','Needs review' if example else 'Not reported',SOURCE if example else '',None])
    ws=sheet(wb,'Resources','Contractor resources','Whole-number counts. Negative variance means actual resources below plan.',['Resource','Planned','Actual','Calculated actual-plan','Printed variance','Unit','Entry status','Source reference','Notes'],rows)
    numeric(ws,'B4:C5');dropdown(ws,'G4:G5',STATUS)
    ws=sheet(wb,'Scope_Layers','Scope quantities and pavement layers','Optional sections: type Scope or Layer. Layer thickness is separate from scope quantity. Blank means not reported.',['Type','Stable ID / layer code','Description','Quantity','Unit','Thickness mm','Qualifier / notes','Entry status','Source reference'],[
      ['Scope' if example else None,'scope_01','Mughsayl Bridge' if example else None,540 if example else None,'m',None,None,'Needs review' if example else 'Not reported',SOURCE if example else ''],
      ['Scope' if example else None,'scope_02','Marnif cave access road' if example else None,2 if example else None,'km',None,'Approximate' if example else None,'Needs review' if example else 'Not reported',SOURCE if example else ''],
      *[[None,f'item_{i:02}',None,None,None,None,None,'Not reported',None] for i in range(3,16)]])
    dropdown(ws,'A4:A18','"Scope,Layer"');dropdown(ws,'H4:H18',STATUS);numeric(ws,'D4:D18');numeric(ws,'F4:F18')
    updates=['Box culvert 4 cell (3x3) completed in all respects.','Animal crossing works including backfilling completed. Approach slabs completed.','Backfilling at parking area in progress.','Precast planks casting in progress – 3500 numbers completed.','Arches piles in progress (76/76).','Deck slab casted from ABT01 to ABT02.','Bridge barriers completed.','Bridge light poles installed.','Bridge asphalt works completed.','Café superstructure works in progress.']
    ws=sheet(wb,'Status_Updates','Narrative progress and independent trade status','Preserve the source wording. For example, “in progress (76/76)” needs clarification, not automatic conversion to Complete.',['Update ID','Description / source wording','Explicit trade status','Entry status','Source reference','Clarification / internal note'],[[f'update_{i+1:02}',updates[i] if example and i<len(updates) else None,'Unknown','Needs review' if example and i<len(updates) else 'Not reported',SOURCE if example and i<len(updates) else None,None] for i in range(20)])
    ws.column_dimensions['B'].width=72;dropdown(ws,'C4:C23','"Complete,Active,Behind,Not started,Unknown"');dropdown(ws,'D4:D23',STATUS)
    parties=[('Client','Ministry of Transport & Communications','Directorate General of Roads & Land Transport'),('Consultant','VIA International Engineering Consultancy',None),('Contractor','AZ Engineers & Partners LLC',None),('Client representative',None,None),('Other contact',None,None)]
    ws=sheet(wb,'Contacts','Project parties and contacts','Enter telephone numbers as TEXT. Unclear small-print names and numbers were not guessed in the example.',['Party role','Organisation','Department','Representative','Position / title','Telephone','Email','Entry status','Source reference'],[[role,org if example else None,dep if example else None,None,None,None,None,'Needs review' if example and org else 'Not reported',SOURCE if example else None] for role,org,dep in parties])
    dropdown(ws,'H4:H8',STATUS)
    for r in range(4,9):ws.cell(r,6).number_format='@'
    assets=[('logo_main','Main dashboard logo','Required',None),('layout_main','Location / project layout','Required','Project Layout on Google Image'),('photo_1','Construction photo 1','Required','Bridge Opening'),('photo_2','Construction photo 2','Required','Bridge Opening'),('photo_3','Construction photo 3','Required','Slope Protection Works'),('photo_4','Construction photo 4','Required','Café Sub Structure'),('logo_client','Additional client logo','Optional',None),('logo_contractor','Additional contractor logo','Optional',None),('status_drawing','Activity-status drawing','If reported','Status of key project activities'),('s_curve','Original S-curve image','If chart table unreadable','Project Physical Progress Chart/S-Curve')]
    ws=sheet(wb,'Assets','Manual image uploads and captions','Upload files separately in the app. Match using Asset ID, not filename guesses. Images are not embedded in this workbook.',['Asset ID','Purpose','Requirement','Local filename','Caption','Secondary description','Capture date','Entry status','Upload guidance'],[[key,purpose,req,None,caption if example else None,None,None,'Needs review' if example and caption else 'Not reported','JPG/PNG/WebP. Keep original resolution; do not upscale to imply extra detail.'] for key,purpose,req,caption in assets])
    ws.column_dimensions['E'].width=40;ws.column_dimensions['I'].width=65;dropdown(ws,'H4:H13',STATUS)
    for r in range(4,14):ws.cell(r,7).number_format='yyyy-mm-dd'
    # Calculated comparisons are explicitly separate from source figures.
    def diff(a,b):return f'=IF(COUNT({refs[a]},{refs[b]})=2,{refs[a]}-{refs[b]},"Not reported")'
    checks=[
      ['Physical actual minus planned',diff('physical_actual','physical_planned'),'Percentage points; compare with the printed physical variance.'],
      ['Financial actual minus planned',diff('financial_actual','financial_planned'),'Sample result -4.39; source prints +4.39. Confirm convention.'],
      ['Elapsed + remaining - construction',f'=IF(COUNT({refs["elapsed_days"]},{refs["remaining_days"]},{refs["construction_days"]})=3,{refs["elapsed_days"]}+{refs["remaining_days"]}-{refs["construction_days"]},"Not reported")','Zero means reported durations reconcile. Sample: 636 + 33 = 669.'],
      ['Original including contingency minus components',f'=IF(COUNT({refs["original_inc_contingency"]},{refs["original_ex_contingency"]},{refs["contingency"]})=3,{refs["original_inc_contingency"]}-{refs["original_ex_contingency"]}-{refs["contingency"]},"Not reported")','Zero means monetary components reconcile.'],
      ['Days between original start and completion',diff('completion_date','start_date'),'Calendar subtraction, excluding the starting day. Compare with contractual counting rules, not an automatic error.'],
      ['Source schedule readability','MANUAL CHECK','Tiny chart table in sample was not transcribed. Supply original data or source image.'],
      ['Project identity and reporting month','MANUAL CHECK','Confirm uploaded PDF, workbook and selected existing project all match.'],
      ['Images and captions','MANUAL CHECK','Check main logo, layout, all four photos and any supplemental drawings.'],
      ['Publication approval','NOT APPROVED','Preview and approve inside the app after resolving all required fields and review issues.'],
    ]
    ws=sheet(wb,'Checks','Review checklist and calculated comparisons','Checks are not a certification of accuracy. Formula results calculate when opened in Excel.',['Check','Result','Reviewer guidance'],checks)
    ws.column_dimensions['A'].width=48;ws.column_dimensions['B'].width=28;ws.column_dimensions['C'].width=100
    for sh in wb:
        for row in sh:
            for c in row:
                if c.data_type=='f':c.fill=PatternFill('solid',fgColor=GREEN);c.number_format='0.000'
    wb.active=0
    name='VIA_Mughsayl_July_2026_Example.xlsx' if example else 'VIA_Project_Report_Template_v1.xlsx'
    OUT.mkdir(exist_ok=True);path=OUT/name;wb.save(path)
    reopened=load_workbook(path)
    assert len(reopened.sheetnames)==12
    assert reopened['Activities']['F4'].data_type=='f'
    assert reopened['Project']['C4'].value == (PROJECT[0][5] if example else None)
    assert reopened['Monthly_Schedule']['B4'].value is None
    assert not any(c.data_type=='e' for s in reopened for row in s for c in row)
    print(f'{path} | {len(reopened.sheetnames)} sheets | verified structure')

if __name__=='__main__':
    build(False);build(True)
